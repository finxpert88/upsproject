from pathlib import Path
from pypdf import PdfReader
import openpyxl,json,csv,pdfplumber,re
p=Path(__file__).resolve().parent/'tmp'/'verified'
expect=json.loads((p/'expected.json').read_text(encoding='utf-8'))
wb=openpyxl.load_workbook(p/'REP-01.xlsx',data_only=False)
han=re.compile('[\u3400-\u9fff]')
required={'REP-01.pdf','REP-01.xlsx','REP-02.xlsx','REP-02.csv','REP-03.pdf','REP-03.xlsx','REP-03.csv','REP-04.pdf','REP-04.xlsx'}
assert all((p/f).is_file() for f in required)
assert wb.sheetnames==['Notes','Metric Dictionary','Data Quality','Current Health','Operating Summary','Details','Alarms & Events'],wb.sheetnames
assert any(row[0].value=='Source' and row[1].value=='Simulated Data / simulation' for row in wb['Notes'])
assert wb['Details'].max_row==len(expect['details'])+1
assert wb['Operating Summary']['C2'].value==expect['summary'][0][2]
assert wb['Operating Summary']['C2'].data_type=='n'
assert wb['Details'].freeze_panes=='A2'
assert wb['Details']['B2'].is_date
assert any(row[0].value=='Generation' and row[1].value==expect['generationId'] for row in wb['Notes'])
assert any(row[0].value=='Time Zone' and row[1].value=='Asia/Singapore' for row in wb['Notes'])
results=[]
for file in sorted(p.glob('*.xlsx')):
 book=openpyxl.load_workbook(file,data_only=False)
 assert all(not han.search(name) for name in book.sheetnames)
 assert any('Simulated Data' in str(c.value) for sheet in book for row in sheet for c in row)
 for sheet in book:
  assert sheet.freeze_panes=='A2'
  for row in sheet:
   for cell in row:
    assert cell.data_type!='f'
    assert not han.search(str(cell.value)),(file.name,sheet.title,cell.coordinate)
  assert all((dim.height or 0)<=409 for dim in sheet.row_dimensions.values())
 results.append({'file':file.name,'sheets':book.sheetnames})
 print(file.name,len(book.sheetnames),'English worksheets; typed values and formula-injection guard passed')
for pdf in p.glob('*.pdf'):
 reader=PdfReader(pdf);texts=[page.extract_text() for page in reader.pages]
 assert all('Simulated Data' in text for text in texts)
 assert all('Page ' in text and '[Asia/Singapore]' in text for text in texts)
 assert all(not han.search(text) for text in texts)
 assert all(len(page['/Resources']['/Font']['/F1']['/DescendantFonts'][0].get_object()['/FontDescriptor']['/FontFile2'].get_data())>10000 for page in reader.pages)
 with pdfplumber.open(pdf) as doc:
  for index,page in enumerate(doc.pages):
   assert all(c['x0']>=20 and c['x1']<=825 and c['top']>=10 and c['bottom']<=590 for c in page.chars), (pdf.name,index,'clipped')
   page.to_image(resolution=110).save(p/(pdf.stem+'-'+str(index+1)+'.png'))
 results.append({'file':pdf.name,'pages':len(reader.pages)})
 print(pdf.name,len(reader.pages),'pages; English text, embedded font, time zone, and text bounds passed')
for file in sorted(p.glob('*.csv')):
 rows=list(csv.reader(file.open(encoding='utf-8-sig',newline='')))
 assert len(rows)>4 and rows[0][1]=='Simulated Data'
 assert all(not han.search(cell) for row in rows for cell in row)
 assert any('Asia/Singapore' in row for row in rows)
 if file.name=='REP-02.csv':assert len(rows)==len(expect['details'])+4
 results.append({'file':file.name,'rows':len(rows)})
 print(file.name,len(rows),'rows; English headers, simulation marker, and time zone passed')
(p/'files-r9.json').write_text(json.dumps(results,indent=2),encoding='utf-8')
print('LANG-03 parsing passed for all 9 allowed outputs; inspect rendered pages/sheets separately.')
